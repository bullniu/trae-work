"""验证 test_filled.docx 的填充效果：打印每个表格去重单元格的内容"""
from docx import Document
from docx.oxml.ns import qn

doc = Document("/workspace/test_filled.docx")

def unique_cells(row):
    seen = set(); out = []
    for c in row.cells:
        cid = id(c._tc)
        if cid not in seen:
            seen.add(cid); out.append(c)
    return out

for ti, table in enumerate(doc.tables):
    print(f"\n{'='*60}\n表{ti} (rows={len(table.rows)})\n{'='*60}")
    for ri, row in enumerate(table.rows):
        cells = unique_cells(row)
        line = []
        for c in cells:
            t = "\n".join(p.text for p in c.paragraphs).strip()
            t = t.replace("\n","/")[:30]
            line.append(t if t else "·空·")
        # 只打印有内容的行或前几行
        if any(c.strip() not in ('·空·','') for c in line) or ri < 3:
            print(f"行{ri}: {' | '.join(line[:8])}")
